#!/usr/bin/env python3
"""Genera SW Test Plan, TC Design y Test Report (plantillas Excel del maestro) desde:

    tests/test_*.py      los TC automatizados: decorador @tc + docstring de cada prueba (tests/tc.py)
    docs/testing/        meta.yml (equipo y proyecto) y manuales.yml (TC que no son código)
    issues de GitHub     los bugs con label «bug» (plantilla «Defecto») -> hoja «Bug report»
    reports/             results.json (pytest) y coverage.json (opcional)

Uso:
    pytest --cov=mis_trapitos --cov-report=json:reports/coverage.json
    python tools/generar_documentos.py            # escribe en entregables/

Los bugs se leen con GITHUB_TOKEN/GH_TOKEN o con la sesión de `gh auth login`.
"""
from __future__ import annotations

import argparse
import ast
import copy
import datetime as dt
import inspect
import io
import json
import math
import os
import re
import subprocess
import sys
import urllib.request
from collections import Counter, OrderedDict
from pathlib import Path

import openpyxl
import yaml
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
TPL = ROOT / "docs" / "testing" / "plantillas"
TESTING = ROOT / "docs" / "testing"
TESTS = ROOT / "tests"

# --------------------------------------------------------------------------- TC desde el código
SECCIONES = {"Precondiciones": "pre", "Pasos": "steps", "Esperado": "expected", "Comentarios": "comments"}
_SECCION = re.compile(r"^(%s):\s*(.*)$" % "|".join(SECCIONES))


def leer_docstring(texto):
    """Docstring del TC -> objective, pre, steps, expected, comments (formato en tests/tc.py)."""
    datos = {"objective": [], "pre": [], "steps": [], "expected": [], "comments": []}
    actual = "objective"
    for linea in inspect.cleandoc(texto or "").split("\n"):
        m = _SECCION.match(linea.strip())
        if m:
            actual = SECCIONES[m.group(1)]
            if m.group(2):
                datos[actual].append(m.group(2))
            continue
        linea = linea.strip()
        if actual in ("steps", "expected"):
            if linea.startswith("- "):
                datos[actual].append(linea[2:])
            elif linea and datos[actual]:
                datos[actual][-1] += " " + linea
        elif linea:
            datos[actual].append(linea)
    for k in ("objective", "pre", "comments"):
        datos[k] = " ".join(datos[k])
    return datos


def _decorador_tc(nodo):
    for d in nodo.decorator_list:
        if isinstance(d, ast.Call) and getattr(d.func, "id", getattr(d.func, "attr", None)) == "tc":
            return d
    return None


def pruebas_del_codigo(tests: Path = TESTS):
    """[(nodeid, nodo_ast, decorador @tc o None, decorador de su clase o None)] de cada prueba."""
    pruebas = []
    for f in sorted(tests.glob("test_*.py")):
        rel = f"tests/{f.name}"
        for n in ast.parse(f.read_text(encoding="utf-8")).body:
            if isinstance(n, ast.FunctionDef) and n.name.startswith("test"):
                pruebas.append((f"{rel}::{n.name}", n, _decorador_tc(n), None))
            elif isinstance(n, ast.ClassDef):
                deco_clase = _decorador_tc(n)
                if deco_clase:
                    pruebas.append((f"{rel}::{n.name}", n, deco_clase, None))
                for s in n.body:
                    if isinstance(s, ast.FunctionDef) and s.name.startswith("test"):
                        pruebas.append((f"{rel}::{n.name}::{s.name}", s, _decorador_tc(s), deco_clase))
    return pruebas


def leer_tc_del_codigo(tests: Path = TESTS):
    """Reúne los TC declarados con @tc. Regresa (casos, errores)."""
    casos, errores = OrderedDict(), []
    for nodeid, nodo, deco, deco_clase in pruebas_del_codigo(tests):
        if deco is None:
            continue
        try:
            tc_id = ast.literal_eval(deco.args[0])
            kw = {k.arg: ast.literal_eval(k.value) for k in deco.keywords}
        except (IndexError, ValueError):
            errores.append(f"{nodeid}: @tc necesita el id como texto y argumentos literales")
            continue
        caso = casos.setdefault(tc_id, {"id": tc_id, "tests": []})
        caso["tests"].append(nodeid)
        if kw:
            if "name" in caso:
                errores.append(f"{tc_id}: los datos del TC están en dos pruebas ({caso['definido_en']} y {nodeid})")
                continue
            caso.update(level=kw.get("nivel"), module=kw.get("modulo"), reqs=list(kw.get("reqs", [])),
                        name=kw.get("nombre"), definido_en=nodeid, **leer_docstring(ast.get_docstring(nodo)))
    for caso in casos.values():
        if "name" not in caso:
            errores.append(f"{caso['id']}: ninguna prueba trae los datos (nivel, modulo, reqs, nombre)")
            continue
        faltan = [k for k in ("level", "module", "name", "reqs", "objective", "pre", "steps", "expected") if not caso.get(k)]
        if faltan:
            errores.append(f"{caso['id']} ({caso['definido_en']}): falta {', '.join(faltan)}")
    return [c for c in casos.values() if "name" in c], errores


def cargar_catalogo(base: Path = TESTING, tests: Path = TESTS):
    """meta.yml + TC del código (@tc) + TC manuales (manuales.yml). Agrega 'errores' de consistencia."""
    leer = lambda f: yaml.safe_load(f.read_text(encoding="utf-8"))
    meta = leer(base / "meta.yml")
    automaticos, errores = leer_tc_del_codigo(tests)
    manuales = leer(base / "manuales.yml") or []
    orden = list(meta["levels"])
    for c in automaticos + manuales:
        if c.get("level") not in orden:
            errores.append(f"{c['id']}: nivel {c.get('level')!r} no está en meta.yml (levels)")
            c["level"] = orden[0]
    ids = [c["id"] for c in automaticos + manuales]
    errores += [f"id duplicado: {i}" for i, n in Counter(ids).items() if n > 1]
    cases = sorted(automaticos + manuales, key=lambda c: (orden.index(c["level"]), c["id"]))
    return {"meta": meta, "cases": cases, "defects": [], "errores": errores}


def coincide(nodeid, prefijo):
    """El resultado de pytest pertenece a la prueba (incluye parámetros y métodos de una clase)."""
    return nodeid == prefijo or nodeid.startswith(prefijo + "[") or nodeid.startswith(prefijo + "::")

# --------------------------------------------------------------------------- bugs desde GitHub
CAMPOS_BUG = {"TC que lo detectó": "tc", "Pasos para reproducir": "pasos", "Severidad": "severidad",
              "Esperado vs. real": "esperado"}


def token_github():
    for var in ("GITHUB_TOKEN", "GH_TOKEN"):
        if os.environ.get(var):
            return os.environ[var]
    try:
        return subprocess.run(["gh", "auth", "token"], capture_output=True, text=True, check=True).stdout.strip() or None
    except (OSError, subprocess.CalledProcessError):
        return None


def leer_formulario(cuerpo):
    """Cuerpo de un issue creado con la plantilla «Defecto» -> {tc, pasos, severidad, esperado}."""
    datos, actual = {}, None
    for linea in (cuerpo or "").replace("\r\n", "\n").split("\n"):
        if linea.startswith("### "):
            actual = CAMPOS_BUG.get(linea[4:].strip())
            continue
        if actual:
            datos[actual] = (datos.get(actual, "") + "\n" + linea).strip()
    return {k: ("" if v == "_No response_" else v) for k, v in datos.items()}


def defecto_desde_issue(issue, nombres):
    f = leer_formulario(issue.get("body"))
    nombre = lambda login: nombres.get(login, login)
    pasos = f.get("pasos", "")
    if f.get("esperado"):
        pasos += f"\n\nEsperado vs. real:\n{f['esperado']}"
    return {"id": f"#{issue['number']}", "tc": f.get("tc", "").strip() or "N/A",
            "description": re.sub(r"^\s*(\[DEF\]|Bug:)\s*", "", issue["title"]),
            "steps": pasos.strip() or "Ver el issue en GitHub.",
            "severity": f.get("severidad") or "Sin clasificar",
            "state": "Closed" if issue["state"] == "closed" else "Open",
            "found_by": nombre(issue["user"]["login"]),
            "assigned_to": ", ".join(nombre(a["login"]) for a in issue.get("assignees", [])) or "Sin asignar"}


def cargar_bugs(meta, token):
    """Issues con label «bug» del repo, ordenados por número."""
    repo = os.environ.get("GITHUB_REPOSITORY") or meta["repo"]
    issues, pagina = [], 1
    while True:
        req = urllib.request.Request(
            f"https://api.github.com/repos/{repo}/issues?labels=bug&state=all&per_page=100&page={pagina}",
            headers={"Authorization": f"Bearer {token}", "Accept": "application/vnd.github+json"})
        with urllib.request.urlopen(req, timeout=30) as r:
            lote = json.load(r)
        issues += [i for i in lote if "pull_request" not in i]
        if len(lote) < 100:
            break
        pagina += 1
    nombres = meta.get("nombres") or {}
    return [defecto_desde_issue(i, nombres) for i in sorted(issues, key=lambda i: i["number"])]

# --------------------------------------------------------------------------- datos
def cargar(results_path: Path, cov_path: Path, con_github=True):
    cat = cargar_catalogo()
    if con_github:
        token = token_github()
        if not token:
            raise SystemExit("Sin acceso a GitHub para leer los bugs: corre `gh auth login`, define GITHUB_TOKEN "
                             "o usa --sin-github (el Bug report saldrá vacío).")
        cat["defects"] = cargar_bugs(cat["meta"], token)
    results = json.loads(results_path.read_text(encoding="utf-8")) if results_path.exists() else {}
    cov = None
    if cov_path.exists():
        cov = json.loads(cov_path.read_text(encoding="utf-8"))
    fecha = dt.date.fromtimestamp(results_path.stat().st_mtime) if results_path.exists() else dt.date.today()
    return cat, results, cov, fecha


def evaluar(cat, results, fecha_ejec):
    """Agrega a cada TC: status (PASSED/FAILED/BLOCKED/NOT_EXECUTED), n_tests, fecha, tester."""
    meta, testers = cat["meta"], cat["meta"]["testers"]
    modulos = list(OrderedDict.fromkeys(c["module"] for c in cat["cases"]))
    defectos = {d["id"]: d for d in cat.get("defects", [])}
    for c in cat["cases"]:
        c["method"] = "TA" if c.get("tests") else "Manual"
        c["tester"] = c.get("executed_by") or testers[modulos.index(c["module"]) % len(testers)]
        hits = [(k, v) for k, v in results.items() if any(coincide(k, p) for p in c.get("tests", []))]
        c["hits"] = hits
        if c["method"] == "TA":
            if not hits:
                c["status"], c["date"] = "NOT_EXECUTED", dt.date.fromisoformat(meta["planned_date"])
            else:
                bad = [h for h in hits if h[1]["outcome"] in ("failed", "xfailed", "xpassed", "error")]
                skipped = [h for h in hits if h[1]["outcome"] == "skipped"]
                c["status"] = "FAILED" if bad else "NOT_EXECUTED" if skipped and len(skipped) == len(hits) else "PASSED"
                c["date"] = fecha_ejec
        else:
            manual = c.get("manual_result")
            c["status"] = "BLOCKED" if c.get("blocked_by") else manual or "NOT_EXECUTED"
            c["date"] = dt.date.fromisoformat(str(c.get("executed_on") or meta["planned_date"]))
        c["n_tests"] = len(hits)
        propios = sorted((d for d in defectos.values() if d["tc"] == c["id"]), key=lambda d: d["state"] != "Open")
        c["defect"] = propios[0] if propios else None
    return cat


def validar(cat, con_github=True):
    errs = list(cat.get("errores", []))
    for lv in cat["meta"]["levels"]:
        n = sum(c["level"] == lv for c in cat["cases"])
        if n < 10:
            errs.append(f"nivel {lv}: {n} TC (<10)")
    cubiertos = {r for c in cat["cases"] for r in c["reqs"]}
    errs += [f"requisito sin TC: {r}" for r in cat["meta"]["requirements"] if r not in cubiertos]
    for c in cat["cases"]:
        if con_github and c["status"] in ("FAILED", "BLOCKED") and not c["defect"] and not any(d["id"] == c.get("blocked_by") for d in cat["defects"]):
            errs.append(f"{c['id']} {c['status']} sin defecto registrado")
        if c["method"] == "TA" and not c["hits"]:
            errs.append(f"{c['id']}: ningún test de pytest coincide con {c['tests']} (¿corriste pytest?)")
    return errs


def conteos(cat):
    cs = cat["cases"]
    st = Counter(c["status"] for c in cs)
    ejec = st["PASSED"] + st["FAILED"]
    por_modulo = Counter(c["module"] for c in cs)
    tests_modulo = Counter()
    tests_res = Counter()
    for c in cs:
        for _, v in c["hits"]:
            tests_modulo[c["module"]] += 1
            tests_res["passed" if v["outcome"] == "passed" else "known_defect" if v["outcome"] == "xfailed" else v["outcome"]] += 1
    sev1 = sum(1 for d in cat["defects"] if d["severity"].startswith("1") and d["state"] != "Closed")
    return dict(total=len(cs), planned=len(cs), executed=ejec, status=st, por_modulo=por_modulo,
                tests_modulo=tests_modulo, tests_res=tests_res, n_tests=sum(tests_modulo.values()),
                pass_pct=(st["PASSED"] / ejec if ejec else 0), critical=sev1)

# --------------------------------------------------------------------------- gráficas
PAL = {"PASSED": "#2E7D32", "FAILED": "#C62828", "BLOCKED": "#EF6C00", "NOT_EXECUTED": "#9E9E9E"}
MODCOL = ["#1F4E79", "#2E86AB", "#F6AE2D", "#8E6C8A", "#5C8A4D", "#B5543C"]


def graficas(cat, k, cov, outdir: Path):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    outdir.mkdir(parents=True, exist_ok=True)
    pngs = {}

    def pie(name, title, labels, values, colors):
        pares = [(l, v, c) for l, v, c in zip(labels, values, colors) if v]
        fig, ax = plt.subplots(figsize=(5.4, 4.2), dpi=130)
        w, _, at = ax.pie([p[1] for p in pares], colors=[p[2] for p in pares], startangle=90, counterclock=False,
                          autopct=lambda p: f"{p:.0f}%" if p >= 4 else "", pctdistance=0.78,
                          wedgeprops=dict(width=0.45, edgecolor="white", linewidth=1.5),
                          textprops=dict(color="white", fontsize=9, fontweight="bold"))
        ax.legend(w, [f"{l}: {v}" for l, v, _ in pares], loc="center left", bbox_to_anchor=(1.0, 0.5), frameon=False, fontsize=9)
        ax.text(0, 0, str(sum(p[1] for p in pares)), ha="center", va="center", fontsize=20, fontweight="bold")
        ax.set_title(title, fontsize=11, loc="left", fontweight="bold")
        fig.tight_layout()
        p = outdir / f"{name}.png"
        fig.savefig(p)
        plt.close(fig)
        pngs[name] = p

    etiq = {"PASSED": "Passed", "FAILED": "Failed", "BLOCKED": "Blocked", "NOT_EXECUTED": "Not executed"}
    pie("pie_total_tc", "Casos de prueba por estado (total)", list(etiq.values()),
        [k["status"][s] for s in etiq], [PAL[s] for s in etiq])
    mods = list(k["por_modulo"])
    pie("pie_tc_modulo", "Casos de prueba por módulo", mods, [k["por_modulo"][m] for m in mods], MODCOL)
    tm = list(k["tests_modulo"])
    pie("pie_pytest_modulo", "Pruebas pytest ejecutadas por módulo", tm, [k["tests_modulo"][m] for m in tm], MODCOL)
    niveles = cat["meta"]["levels"]
    pie("pie_tc_nivel", "Casos de prueba por nivel", list(niveles.values()),
        [sum(c["level"] == l for c in cat["cases"]) for l in niveles], MODCOL)
    # barras apiladas por módulo y estado
    fig, ax = plt.subplots(figsize=(6.4, 3.6), dpi=130)
    bottom = [0] * len(mods)
    for s in etiq:
        vals = [sum(1 for c in cat["cases"] if c["module"] == m and c["status"] == s) for m in mods]
        ax.bar(mods, vals, bottom=bottom, color=PAL[s], label=etiq[s], edgecolor="white")
        bottom = [a + b for a, b in zip(bottom, vals)]
    ax.set_title("Estado de los casos de prueba por módulo", fontsize=11, loc="left", fontweight="bold")
    ax.legend(frameon=False, fontsize=8, ncol=4, loc="upper center", bbox_to_anchor=(0.5, -0.12))
    ax.spines[["top", "right"]].set_visible(False)
    fig.tight_layout()
    p = outdir / "bar_estado_modulo.png"
    fig.savefig(p)
    plt.close(fig)
    pngs["bar_estado_modulo"] = p
    if cov:
        files = sorted(((f.replace("mis_trapitos/", ""), d["summary"]["percent_covered"]) for f, d in cov["files"].items()),
                       key=lambda x: x[1])
        fig, ax = plt.subplots(figsize=(6.4, max(3, 0.28 * len(files) + 1)), dpi=130)
        ax.barh([f for f, _ in files], [v for _, v in files], color="#1F4E79")
        ax.axvline(cov["totals"]["percent_covered"], color="#C62828", ls="--", lw=1)
        ax.set_xlim(0, 100)
        ax.set_title(f"Cobertura de líneas por archivo (total {cov['totals']['percent_covered']:.1f}%)",
                     fontsize=11, loc="left", fontweight="bold")
        ax.tick_params(axis="y", labelsize=7)
        ax.spines[["top", "right"]].set_visible(False)
        fig.tight_layout()
        p = outdir / "bar_cobertura.png"
        fig.savefig(p)
        plt.close(fig)
        pngs["bar_cobertura"] = p
    return pngs

# --------------------------------------------------------------------------- utilidades Excel
THIN = Side(style="thin", color="999999")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def lineas(texto, ancho):
    return sum(max(1, math.ceil(len(l) / ancho)) for l in str(texto).split("\n"))


def alto(texto, ancho, base=13.0):
    return max(15.0, lineas(texto, ancho) * base + 4)


def copiar_estilo(src, dst):
    dst.font, dst.border, dst.fill = copy.copy(src.font), copy.copy(src.border), copy.copy(src.fill)
    dst.alignment, dst.number_format, dst.protection = copy.copy(src.alignment), src.number_format, copy.copy(src.protection)


def extraer_imagenes(ws):
    """Lee una sola vez los bytes de las imágenes (openpyxl cierra el stream tras _data())."""
    return [(im._data(), im.width, im.height, copy.deepcopy(im.anchor)) for im in ws._images]


def copiar_imagenes(imgs, dst_ws):
    for data, w, h, anchor in imgs:
        n = XLImage(io.BytesIO(data))
        n.width, n.height, n.anchor = w, h, copy.deepcopy(anchor)
        dst_ws.add_image(n)


def copiar_hoja(wb, master, titulo, imgs):
    ws = wb.copy_worksheet(master)
    ws.title = titulo
    copiar_imagenes(imgs, ws)
    ws.sheet_view.showGridLines = master.sheet_view.showGridLines
    for dv in master.data_validations.dataValidation:
        ws.add_data_validation(DataValidation(type=dv.type, formula1=dv.formula1, sqref=str(dv.sqref), allow_blank=True))
    return ws


def ajustar_impresion(wb):
    """Cada hoja cabe a lo ancho de una página (útil para imprimir/PDF)."""
    from openpyxl.worksheet.properties import PageSetupProperties
    for ws in wb.worksheets:
        ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        ws.page_setup.orientation = "landscape" if ws.title in ("Matrix", "Detailed status", "Bug report", "Strategy") else "portrait"


def numerar(items):
    return "\n".join(f"{i}.- {t}" for i, t in enumerate(items, 1))


def poner(ws, ref, valor, wrap=True, h=None):
    c = ws[ref]
    c.value = valor
    if wrap:
        a = copy.copy(c.alignment)
        c.alignment = Alignment(horizontal=a.horizontal, vertical=a.vertical or "top", wrap_text=True)
    return c

# --------------------------------------------------------------------------- SW Test Plan
def construir_plan(cat, k, fecha, out: Path, tpl: Path):
    meta = cat["meta"]
    wb = openpyxl.load_workbook(tpl)
    master = wb["UTS_TestCase_1"]
    imgs_master = extraer_imagenes(master)
    # ---- intro
    wi = wb["intro"]
    for col, v in (("C", fecha.isoformat()), ("D", meta["department"]),
                   ("E", "v1.0 — Primera emisión: alcance, estrategia, matriz de trazabilidad y diseño de casos de prueba "
                         f"({k['total']} TC en {len(meta['levels'])} niveles) generados desde las pruebas de tests/ y los issues de GitHub.")):
        c = wi[f"{col}14"]
        c.value = v
        c.alignment = Alignment(wrap_text=True, vertical="top")
    wi.row_dimensions[14].height = 52
    for fila, rev in enumerate(meta.get("revisiones") or [], start=15):   # revisión y aprobación en equipo
        for col, v in (("C", str(rev["fecha"])), ("D", rev["quien"]), ("E", rev["descripcion"])):
            c = wi[f"{col}{fila}"]
            c.value = v
            c.alignment = Alignment(wrap_text=True, vertical="top")
        wi.row_dimensions[fila].height = 30
    wi["B6"].value = f"Proyecto: {meta['project_title']} (versión {meta['release']})"
    # ---- Matrix
    wm = wb["Matrix"]
    wm["C1"].value = "SDD.md, SAD.md, tests/ (@tc), docs/testing/manuales.yml"
    wm["C2"].value = f"{meta['release']} — Test Plan v1.0 ({fecha.isoformat()})"
    for rng in list(wm.merged_cells.ranges):
        if str(rng) == "B8:G8":
            wm.unmerge_cells(str(rng))
    modelo = {col: wm[f"{col}7"] for col in "BCDEFG"}
    estilos = {col: copy.copy(modelo[col]._style) for col in "BCDEFG"}
    for r in range(6, 12):
        for col in "BCDEFG":
            wm[f"{col}{r}"].value = None
    dv_old = list(wm.data_validations.dataValidation)
    wm.data_validations.dataValidation = []
    for i, c in enumerate(cat["cases"]):
        r = 6 + i
        for col in "BCDEFG":
            wm[f"{col}{r}"]._style = copy.copy(estilos[col])
        wm.row_dimensions[r].height = 30
        c["matrix_row"] = r
        wm[f"B{r}"] = c["id"]
        wm[f"C{r}"] = f"={c['id']}!$B$4"
        wm[f"D{r}"] = c["method"]
        wm[f"E{r}"] = "X"
        wm[f"F{r}"] = f'=IF({c["id"]}!$D$5="Complete",{c["id"]}!$D$17,{c["id"]}!$D$5)'
        nota = []
        if c.get("defect"):
            nota.append(c["defect"]["id"])
        if c.get("blocked_by"):
            nota.append("Bloqueado por " + c["blocked_by"])
        wm[f"G{r}"] = ", ".join(nota)
        wm[f"C{r}"].alignment = Alignment(wrap_text=True, vertical="center")
    from openpyxl.formatting.formatting import ConditionalFormattingList
    from openpyxl.formatting.rule import CellIsRule
    wm.conditional_formatting = ConditionalFormattingList()
    rango = f"F6:F{5 + len(cat['cases'])}"
    for texto, color in (("PASSED", "00FF00"), ("FAILED", "FF0000"), ("Not done", "D9D9D9")):
        wm.conditional_formatting.add(rango, CellIsRule(operator="equal", formula=[f'"{texto}"'],
                                                         fill=PatternFill(start_color=color, end_color=color, fill_type="solid")))
    for r in range(6, 6 + len(cat["cases"])):
        wm[f"F{r}"].alignment = Alignment(horizontal="center", vertical="center")
    dv = DataValidation(type="list", formula1='"TA,Manual"', allow_blank=True)
    wm.add_data_validation(dv)
    dv.add(f"D6:D{5 + len(cat['cases'])}")
    # ---- hojas de TC
    for c in cat["cases"]:
        ws = copiar_hoja(wb, master, c["id"], imgs_master)
        poner(ws, "B4", c["name"])
        ws["B5"].value = c["id"]
        ws["C5"].value = f'=IF(Matrix!$E${c["matrix_row"]}="X","YES","NO")'
        ws["D5"].value = "Complete" if c["status"] in ("PASSED", "FAILED") else "Not done"
        poner(ws, "B7", ", ".join(c["reqs"]))
        poner(ws, "B9", c["objective"])
        env = meta["environment"].strip() + ("\n\nMétodo: automatizado con pytest (CI: .github/workflows/tests.yml)." if c["method"] == "TA"
                                             else "\n\nMétodo: manual, ejecutado por el responsable indicado en el Test Report.")
        poner(ws, "B11", env)
        poner(ws, "B13", numerar(c["steps"]))
        poner(ws, "B15", numerar(c["expected"]))
        if c["method"] == "TA":
            nodos = "\n".join(f"  · {h[0]} → {h[1]['outcome']}" for h in c["hits"][:8])
            if len(c["hits"]) > 8:
                nodos += f"\n  · … (+{len(c['hits']) - 8} más)"
            res = f"Ejecución automática {c['date'].isoformat()}: {c['n_tests']} prueba(s) pytest.\n{nodos}"
        elif c["status"] == "BLOCKED":
            res = f"Bloqueado por {c['blocked_by']}; no ejecutable hasta resolverlo."
        else:
            res = f"Pendiente de ejecución manual (planeada: {c['date'].isoformat()}). Responsable: {c['tester']}."
        pre = f"Precondiciones:\n-{c['pre']}\n\nResultado de ejecución:\n{res}"
        poner(ws, "B17", pre)
        if c["status"] in ("PASSED", "FAILED"):
            ws["D17"].value = c["status"]
        poner(ws, "B19", c.get("comments", "") or ("Caso derivado de " + ", ".join(c["reqs"]) + "."))
        poner(ws, "B21", meta["references"])
        for ref, txt, w in (("B4", c["name"], 55), ("B7", ", ".join(c["reqs"]), 75), ("B9", c["objective"], 80), ("B11", env, 80),
                            ("B13", numerar(c["steps"]), 80), ("B15", numerar(c["expected"]), 80), ("B17", pre, 70),
                            ("B19", ws["B19"].value, 80), ("B21", meta["references"], 80)):
            ws.row_dimensions[ws[ref].row].height = alto(txt, w)
    del wb["UTS_TestCase_1"]
    del wb["Default_page"]
    # ---- hoja Strategy (estrategia + trazabilidad)
    construir_strategy(wb, cat, k, fecha)
    wb.move_sheet("Strategy", offset=-(len(wb.sheetnames) - 2))
    out.parent.mkdir(parents=True, exist_ok=True)
    ajustar_impresion(wb)
    wb.save(out)


def construir_strategy(wb, cat, k, fecha):
    meta = cat["meta"]
    ws = wb.create_sheet("Strategy")
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 2
    ws.column_dimensions["B"].width = 26
    ws.column_dimensions["C"].width = 62
    ws.column_dimensions["D"].width = 52
    H = Font(bold=True, size=14, color="1F4E79")
    S = Font(bold=True, color="FFFFFF")
    FS = PatternFill("solid", fgColor="1F4E79")
    row = [2]

    def titulo(t):
        ws.cell(row[0], 2, t).font = H
        row[0] += 1

    def tabla(cabecera, filas, anchos=(26, 62, 52)):
        for j, h in enumerate(cabecera):
            c = ws.cell(row[0], 2 + j, h)
            c.font, c.fill, c.border = S, FS, BOX
            c.alignment = Alignment(wrap_text=True, vertical="center")
        row[0] += 1
        for f in filas:
            hmax = 15
            for j, v in enumerate(f):
                c = ws.cell(row[0], 2 + j, v)
                c.border, c.alignment = BOX, Alignment(wrap_text=True, vertical="top")
                hmax = max(hmax, alto(v, anchos[j] * 1.05))
            ws.row_dimensions[row[0]].height = hmax
            row[0] += 1
        row[0] += 1

    titulo(f"Estrategia de pruebas — {meta['project_title']} v{meta['release']}")
    tabla(["Sección", "Contenido", "Notas"], [
        ("1. Objetivo", "Verificar que el sistema de punto de venta e inventario «Mis Trapitos» cumple RF-01…RF-26 y RNF-01 del SDD, "
                        "pasando de verificación manual a pruebas automatizadas repetibles en cada Pull Request.",
         "Historias de usuario #1, #4, #26 y #27 del Product Backlog."),
        ("2. Alcance (incluido)", "Capa de datos (conexión y repositorios), capa de servicios (reglas de ventas, devoluciones y cancelaciones), "
                                  "promociones, 12 reportes (RF-15…RF-26), interfaz Tkinter (pantallas, formularios, login) bajo Xvfb, "
                                  "escenarios de aceptación VD-01…VD-09 y conformidad con estándares (ISO 8601, RFC 4180, FIPS 180-4, "
                                  "integridad SQLite) y con los contratos del SDD.",
         "Todas las funcionalidades planteadas desde el inicio del proyecto: ver matriz de trazabilidad abajo."),
        ("3. Fuera de alcance", "Servicio LAN / sincronización entre terminales (B-01), autorización por rol (B-04), migraciones versionadas (B-08).",
         "Quedan registradas como brechas; los TC correspondientes están Bloqueados o Pendientes y con defecto asociado."),
        ("4. Niveles de prueba", "\n".join(f"{v} ({kk}): {sum(c['level'] == kk for c in cat['cases'])} TC" for kk, v in meta["levels"].items()),
         "Mínimo exigido: 10 TC por nivel (verificado automáticamente en tests/test_catalogo.py). Sistema incluye pruebas "
         "funcionales, de regresión, de carga (STS_CAR) y de recuperación (STS_REC); Aceptación incluye los TC manuales "
         "que ejecuta el cliente."),
        ("5. Enfoque", "Automatizado (TA) con pytest para todo lo reproducible; Manual para lo que depende de personas (el cliente) o de varias máquinas. "
                       "Cada TC automatizado se declara en su propia prueba con @tc y su docstring; los manuales en docs/testing/manuales.yml. "
                       "Cada prueba usa una base SQLite temporal (load_examples=False) para no tocar mis_trapitos.db.",
         "El CI genera estos documentos en cada PR con los resultados reales de pytest; un TC manual sin ejecutar se reporta como «Not executed»."),
        ("6. Entorno", meta["environment"].strip(), "Dependencias: requirements.txt / requirements-dev.txt."),
        ("7. Criterios de entrada", "Rama con refactor modular integrada; pytest instalable; BD temporal disponible.", ""),
        ("8. Criterios de salida", "100 % de TC automatizados ejecutados; 0 defectos Críticos abiertos que afecten RF; todos los RF con ≥ 1 TC; "
                                   "cobertura de líneas reportada; defectos abiertos justificados en el Bug report.", ""),
        ("9. Herramientas", "pytest, pytest-cov, GitHub Actions (tests.yml en cada PR), Grafana (seguimiento de métricas), tools/generar_documentos.py.", ""),
        ("9b. Gestión de defectos", "Cada defecto es un issue de GitHub con la plantilla «Defecto» (label bug): TC que lo detectó, pasos, severidad y "
                                    "esperado vs. real. El Bug report se llena desde esos issues; quien lo abre es «Found by» y el assignee «Assigned to».",
         "Un TC Failed o Blocked debe tener su issue de bug; el generador lo verifica."),
        ("10. Roles", "Responsable de calidad: ejecuta suite y publica reporte. Testers: " + ", ".join(meta["testers"]) + ". Desarrollo: corrige defectos.", ""),
        ("11. Riesgos", "R1 Sin sincronización LAN (B-01) → RNF-01 no verificable entre terminales.\nR2 Reglas de reportes ambiguas (B-05).\n"
                        "R3 Pruebas de interfaz dependientes de Xvfb/Tk.\nR4 Fechas dependientes de la hora del sistema.",
         "Mitigación: BD temporal, fechas relativas a hoy, defectos y TC bloqueados documentados."),
        ("12. Calendario", f"Emisión del plan: {fecha.isoformat()}. TC manuales planeados para: {meta['planned_date']}.", ""),
    ])
    titulo("Matriz de trazabilidad requisito → casos de prueba")
    ids_extra = meta.get("other_ids", {})
    filas = []
    for r in meta["requirements"] + list(ids_extra):
        tcs = [c for c in cat["cases"] if r in c["reqs"]]
        filas.append((r if r not in ids_extra else f"{r} — {ids_extra[r]}", ", ".join(c["id"] for c in tcs),
                      " / ".join(f"{s}:{sum(c['status'] == s for c in tcs)}" for s in ("PASSED", "FAILED", "BLOCKED", "NOT_EXECUTED") if any(c['status'] == s for c in tcs))))
    tabla(["Requisito", "Casos de prueba", "Estado (PASSED / FAILED / BLOCKED / NOT_EXECUTED)"], filas)

# --------------------------------------------------------------------------- TC Design (solo hojas TC)
def construir_tc_design(plan_path: Path, out: Path):
    wb = openpyxl.load_workbook(plan_path)
    for n in ("intro", "Matrix", "Strategy"):
        pass
    for ws in wb.worksheets:
        if ws.title not in ("intro", "Matrix", "Strategy"):
            ws["C5"].value = "YES"
    del wb["Matrix"]
    wb["intro"]["B4"].value = "Test Case Design"
    wb["intro"]["B8"].value = "Type of Test Document:  TC Design"
    wb.save(out)

# --------------------------------------------------------------------------- Test Report
ESTADO = {"PASSED": "Pass", "FAILED": "Fail", "BLOCKED": "Blocked", "NOT_EXECUTED": "Not executed"}
FILL = {"Pass": "C8E6C9", "Fail": "FFCDD2", "Blocked": "FFE0B2", "Not executed": "E0E0E0"}


def construir_reporte(cat, k, cov, fecha, pngs, out: Path, tpl: Path):
    meta = cat["meta"]
    wb = openpyxl.load_workbook(tpl)
    s = wb["Summary"]
    s._images = []
    for ref, v in (("C2", meta["project_title"]), ("C3", meta["release"]), ("C4", meta["milestone"]), ("C5", fecha.isoformat())):
        s[ref] = v
        s[ref].font = copy.copy(s["B2"].font)
        s[ref].alignment = Alignment(horizontal="left")
    s["F19"].value = None
    s["D9"], s["D10"], s["D11"] = k["total"], k["planned"], k["executed"]
    s["E10"].value = ("Los TC manuales pendientes y los bloqueados por brechas del SDD no cuentan como ejecutados "
                      f"({k['planned'] - k['executed']} sin ejecutar).")
    pctfmt = s["D15"].number_format
    estilo_c, estilo_d = copy.copy(s["C12"]._style), copy.copy(s["D12"]._style)
    estilo_pct_c, estilo_pct_d = copy.copy(s["C15"]._style), copy.copy(s["D15"]._style)
    estilo_crit_c, estilo_crit_d = copy.copy(s["C16"]._style), copy.copy(s["D16"]._style)
    for r in range(12, 30):
        for col in "BCDEF":
            s[f"{col}{r}"].value = None
    r = 12
    for m, n in k["por_modulo"].items():
        s[f"C{r}"], s[f"D{r}"] = f"     Number of TC {m}", n
        s[f"C{r}"]._style, s[f"D{r}"]._style = copy.copy(estilo_c), copy.copy(estilo_d)
        r += 1
    filas = [("Pass Percentage of the defects (TC passed / executed)", k["pass_pct"], estilo_pct_c, estilo_pct_d, pctfmt),
             ("Number of critical defects (High, abiertos)", k["critical"], estilo_crit_c, estilo_crit_d, "General"),
             ("Total pytest tests executed (passed / known-defect xfail / failed)",
              f"{k['n_tests']} ({k['tests_res']['passed']} / {k['tests_res']['known_defect']} / {k['tests_res']['failed']})", estilo_crit_c, estilo_crit_d, "General")]
    if cov:
        filas.append(("Code coverage (line, pytest-cov)", cov["totals"]["percent_covered"] / 100, estilo_pct_c, estilo_pct_d, "0.0%"))
    for t, v, sc, sd, nf in filas:
        s[f"C{r}"], s[f"D{r}"] = t, v
        s[f"C{r}"]._style, s[f"D{r}"]._style = copy.copy(sc), copy.copy(sd)
        s[f"D{r}"].number_format = nf
        s[f"D{r}"].alignment = Alignment(horizontal="right")
        r += 1
    s.column_dimensions["D"].width = 18
    # diagramas
    fila_img = r + 2
    s[f"B{fila_img - 1}"] = "Diagramas"
    s[f"B{fila_img - 1}"].font = Font(bold=True, size=12, color="1F4E79")
    orden = [("pie_total_tc", "B", 0), ("pie_tc_modulo", "E", 0), ("pie_pytest_modulo", "B", 1), ("pie_tc_nivel", "E", 1),
             ("bar_estado_modulo", "B", 2), ("bar_cobertura", "E", 2)]
    for name, col, lvl in orden:
        if name not in pngs:
            continue
        im = XLImage(str(pngs[name]))
        factor = 470 / im.width
        im.width, im.height = int(im.width * factor), int(im.height * factor)
        s.add_image(im, f"{col}{fila_img + lvl * 20}")
    # ---- Detailed status
    d = wb["Detailed status"]
    modelo = [copy.copy(d.cell(2, c)._style) for c in range(1, 10)]
    for row in d["A2:I40"]:
        for c in row:
            c.value = None
    prev = None
    for i, c in enumerate(cat["cases"]):
        rr = 2 + i
        for j in range(1, 10):
            d.cell(rr, j)._style = copy.copy(modelo[j - 1])
            d.cell(rr, j).alignment = Alignment(wrap_text=True, vertical="top")
        d.cell(rr, 1, c["module"] if c["module"] != prev else None)
        d.cell(rr, 2, f"{c['id']} — {c['name']}")
        d.cell(rr, 3, cat["meta"]["levels"][c["level"]])
        d.cell(rr, 4, c["tester"])
        d.cell(rr, 5, c["date"]).number_format = "yyyy-mm-dd"
        est = ESTADO[c["status"]]
        d.cell(rr, 6, est).fill = PatternFill("solid", fgColor=FILL[est])
        d.cell(rr, 6).font = Font(bold=True, color="000000")
        df = c.get("defect") or next((x for x in cat["defects"] if x["id"] == c.get("blocked_by")), None)
        if df:
            d.cell(rr, 7, f"{df['id']}: {df['description']}")
            d.cell(rr, 8, df["severity"])
            d.cell(rr, 9, df["state"])
        d.row_dimensions[rr].height = alto(f"{c['id']} — {c['name']}", 40)
        prev = c["module"]
    d.column_dimensions["A"].width = 13
    d.column_dimensions["B"].width = 52
    d.column_dimensions["G"].width = 50
    # ---- Bug report
    b = wb["Bug report"]
    modelo = [copy.copy(b.cell(3, c)._style) for c in range(1, 9)]
    for row in b["A3:H40"]:
        for c in row:
            c.value = None
    for i, df in enumerate(cat["defects"]):
        rr = 3 + i
        for j, v in enumerate([df["id"], df["tc"], df["description"], df["steps"], df["severity"], df["state"],
                               df["found_by"], df["assigned_to"]], 1):
            cell = b.cell(rr, j, v)
            cell._style = copy.copy(modelo[j - 1])
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        b.row_dimensions[rr].height = alto(df["steps"], 45)
    for col, w in (("C", 45), ("D", 55), ("H", 24)):
        b.column_dimensions[col].width = w
    out.parent.mkdir(parents=True, exist_ok=True)
    ajustar_impresion(wb)
    wb.save(out)

# --------------------------------------------------------------------------- main
def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--results", default=str(ROOT / "reports" / "results.json"))
    ap.add_argument("--coverage", default=str(ROOT / "reports" / "coverage.json"))
    ap.add_argument("--out", default=str(ROOT / "entregables"))
    ap.add_argument("--sin-github", action="store_true", help="no leer los bugs de GitHub (Bug report vacío)")
    a = ap.parse_args(argv)
    cat, results, cov, fecha = cargar(Path(a.results), Path(a.coverage), con_github=not a.sin_github)
    evaluar(cat, results, fecha)
    errs = validar(cat, con_github=not a.sin_github)
    if errs:
        print("ERRORES DE CONSISTENCIA:\n  - " + "\n  - ".join(errs), file=sys.stderr)
        return 1
    k = conteos(cat)
    out = Path(a.out)
    sufijo = f"{cat['meta']['team']}_{cat['meta']['project']}"
    pngs = graficas(cat, k, cov, out / "diagramas")
    plan = out / f"SW_Test_Plan_{sufijo}.xlsx"
    construir_plan(cat, k, fecha, plan, TPL / "SW_Test_Plan__Team0X___ProjectName_.xlsx")
    construir_tc_design(plan, out / f"TC_Design_{sufijo}.xlsx")
    construir_reporte(cat, k, cov, fecha, pngs, out / f"SWT_summary_report_{sufijo}.xlsx",
                      TPL / "SWT_summary_report__Team0X___ProjectName_.xlsx")
    print(f"TC totales {k['total']} | ejecutados {k['executed']} | Pass {k['status']['PASSED']} | Fail {k['status']['FAILED']} "
          f"| Blocked {k['status']['BLOCKED']} | Not executed {k['status']['NOT_EXECUTED']} | pass% {k['pass_pct']:.0%}")
    print("Archivos en", out)
    return 0


if __name__ == "__main__":
    sys.exit(main())
